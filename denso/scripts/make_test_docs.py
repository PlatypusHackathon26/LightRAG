"""Build DOCX / XLSX / TXT test uploads from facts in the existing DENSO documents."""

import sys
from pathlib import Path

import docx
import openpyxl

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)

# DOCX: SCV guide (English section, pages 3-4) as a Vietnamese workshop sheet.
d = docx.Document()
d.add_heading("Hướng dẫn xưởng – Thay van SCV trên bơm Common Rail", level=1)
d.add_paragraph("Bộ kit DCRS300260 (294009-2150). Tài liệu nội bộ dịch từ hướng dẫn lắp đặt DENSO.")
d.add_heading("Tháo van SCV cũ", level=2)
for step in ["Đánh dấu vị trí giắc nối của SCV so với bơm.",
             "Nới lỏng và tháo các bu-lông.",
             "Tháo gioăng O-ring lớn (vị trí 1)."]:
    d.add_paragraph(step, style="List Number")
d.add_heading("Lắp van SCV mới", level=2)
for step in ["Lắp 1 gioăng O-ring mới.",
             "Lắp 2 chốt dẫn hướng vào lỗ bắt bu-lông và lắp các gioăng mới; bôi dầu động cơ lên O-ring nhỏ.",
             "Đẩy SCV vào thân bơm thật cẩn thận để không làm hỏng O-ring, rồi rút các chốt dẫn hướng ra.",
             "Lắp bu-lông và siết với mô-men xoắn 6,9 đến 10,8 Nm."]:
    d.add_paragraph(step, style="List Number")
d.add_paragraph("QUAN TRỌNG: bắt buộc dùng 2 chốt dẫn hướng khi lắp; lắp không có chốt có thể gây rò rỉ nhiên liệu.")
d.add_heading("Thông số", level=2)
t = d.add_table(rows=1, cols=2)
t.style = "Table Grid"
t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Hạng mục", "Giá trị"
for k, v in [("Mô-men xoắn bu-lông SCV", "6,9 – 10,8 Nm"), ("Số chốt dẫn hướng", "2"),
             ("Kiểm tra sau lắp", "Máy chẩn đoán DENSO DST-PC / DENSO C, không có mã lỗi DTC")]:
    row = t.add_row().cells
    row[0].text, row[1].text = k, v
d.save(out / "Huong_dan_xuong_thay_SCV.docx")

# XLSX: two sheets, oil table (Oils leaflet p.1) and torques (SCV guide p.4, Spark Plug Catalogue).
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Dầu máy nén"
ws.append(["Tên dầu DENSO", "Loại dầu gốc", "Mã DENSO (P/N)", "Dùng cho"])
ws.append(["ND-oil 8", "PAG 46", "DND08250", "Máy nén cơ dùng R134a"])
ws.append(["ND-oil 11", "POE", "DND11250", "Máy nén điện dùng R134a và R1234yf"])
ws.append(["ND-oil 12", "PAG 46 (+ phụ gia)", "DND12250", "Máy nén cơ dùng R134a và R1234yf"])
ws2 = wb.create_sheet("Mô-men xoắn")
ws2.append(["Chi tiết", "Mô-men xoắn", "Ghi chú"])
ws2.append(["Bu-lông van SCV (DCRS300260)", "6,9 – 10,8 Nm", "Dùng 2 chốt dẫn hướng"])
ws2.append(["Bugi M14 mới", "20 – 30 N·m", "Khoảng 1/2 vòng"])
ws2.append(["Bugi M12 mới", "15 – 20 N·m", ""])
wb.save(out / "Bang_tra_nhanh_dau_va_momen.xlsx")

# TXT: run-in procedure and oil formula (AC Compressor Installation Manual p.5), plain UTF-8.
(out / "Ghi_chu_ky_thuat_may_nen_AC.txt").write_text("""GHI CHÚ KỸ THUẬT – LẮP MÁY NÉN ĐIỀU HÒA MỚI

1. Quy trình chạy rà sau khi lắp máy nén mới
Mục đích: đưa dầu máy nén đi khắp hệ thống và bôi trơn ngay từ đầu, tránh hỏng máy nén ngay sau khi lắp.
  1) Đặt nhiệt độ ở mức lạnh nhất.
  2) Bật quạt gió ở tốc độ cao nhất.
  3) Nổ máy và giữ động cơ ở vòng tua không tải.
  4) Bật A/C tối thiểu 5 phút. KHÔNG tăng tốc động cơ!
  5) Sau 5 phút, toàn bộ dầu có sẵn trong máy nén đã đi khắp hệ thống; lúc này mới được tăng tốc và kiểm tra A/C.

2. Điều chỉnh lượng dầu của máy nén mới (Quy trình 1 – không cần súc rửa)
  A = tổng lượng dầu trong máy nén mới
  B = lượng dầu xả ra từ máy nén cũ
  C = lượng dầu cần xả bớt khỏi máy nén mới
  Công thức: A - B = C
  Chỉ dùng đúng loại dầu gốc, không trộn với dầu khác hay dầu đa năng; dùng sai loại dầu sẽ mất bảo hành.
""", encoding="utf-8")
print("\n".join(str(p) for p in sorted(out.iterdir())))
